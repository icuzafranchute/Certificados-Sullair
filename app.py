import streamlit as st
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas as pdf_canvas
from reportlab.platypus import Table, TableStyle
import os, re, io, hashlib, json, zipfile
from datetime import datetime, date
from pathlib import Path

st.set_page_config(page_title="Sullair — Certificados", page_icon="📄", layout="centered")

# ─────────────────────────────────────────────────────────────────────────────
# USUARIOS
# ─────────────────────────────────────────────────────────────────────────────
USUARIOS_FILE = Path(__file__).parent / "usuarios.json"

def _hash(pw): return hashlib.sha256(pw.encode()).hexdigest()

def cargar_usuarios():
    if USUARIOS_FILE.exists():
        with open(USUARIOS_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    usuarios = {"JICUZA": {"nombre":"Javier","apellido":"Icuza",
        "sucursal":"Neuquén","mail":"jicuza@sullair.com","password":_hash("sullair2026")}}
    guardar_usuarios(usuarios)
    return usuarios

def guardar_usuarios(u):
    with open(USUARIOS_FILE, 'w', encoding='utf-8') as f:
        json.dump(u, f, ensure_ascii=False, indent=2)

SUCURSALES = ["Neuquén","Buenos Aires","Mendoza","Bahía Blanca","Comodoro Rivadavia",
              "Salta","Tucumán","San Juan","Córdoba","Mar del Plata","Olavarría"]

UN_MAP = {'1':'Compresor','2':'Generador','3':'Plataforma',
          '4':'Manipulador','5':'Luminaria','6':'Movimiento de Suelo'}

# ─────────────────────────────────────────────────────────────────────────────
# CSS
# ─────────────────────────────────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Syne:wght@700;800&family=DM+Sans:wght@300;400;500;600&display=swap');
html, body, [class*="css"] { font-family: 'DM Sans', sans-serif; }
.titulo { font-family:'Syne',sans-serif; font-size:1.5rem; font-weight:800; color:#00783C; }
.sub    { font-size:.82rem; color:#888; margin-bottom:20px; }
.card   { background:#fff; border:1.5px solid #E0E0E0; border-radius:14px;
          padding:22px 26px; margin-bottom:16px; box-shadow:0 2px 10px rgba(0,0,0,.05); }
.card h3{ font-family:'Syne',sans-serif; color:#00783C; font-size:.95rem; margin-bottom:12px; }
.stat   { background:#E8F5EE; border:1.5px solid #b2dfc5; border-radius:10px;
          padding:12px; text-align:center; }
.stat .n{ font-family:'Syne',sans-serif; font-size:1.8rem; font-weight:800; color:#00783C; display:block; }
.stat .l{ font-size:.72rem; color:#555; }
.stButton>button { background:#00783C!important; color:#fff!important;
    font-family:'Syne',sans-serif!important; font-weight:700!important;
    border-radius:10px!important; border:none!important; width:100%; }
div[data-testid="stDownloadButton"]>button {
    background:#fff!important; color:#00783C!important;
    border:2px solid #00783C!important; font-family:'Syne',sans-serif!important;
    font-weight:700!important; border-radius:10px!important; width:100%; }
div[data-testid="stDownloadButton"]>button:hover { background:#00783C!important; color:#fff!important; }
</style>
""", unsafe_allow_html=True)

# ─────────────────────────────────────────────────────────────────────────────
# LOGIN / REGISTRO
# ─────────────────────────────────────────────────────────────────────────────
def pantalla_auth():
    st.markdown("""
    <style>
    .auth-box { max-width:420px; margin:40px auto 0; background:#fff;
        border-radius:16px; padding:36px 36px 28px; box-shadow:0 4px 24px rgba(0,0,0,.10);
        border:1.5px solid #E0E0E0; }
    .auth-titulo { font-family:'Syne',sans-serif; font-size:1.4rem;
        font-weight:800; color:#00783C; text-align:center; margin-bottom:4px; }
    .auth-sub { font-size:.82rem; color:#888; text-align:center; margin-bottom:20px; }
    </style>""", unsafe_allow_html=True)
    _, col, _ = st.columns([0.5, 3, 0.5])
    with col:
        st.markdown('<div class="auth-box">', unsafe_allow_html=True)
        st.markdown('<p class="auth-titulo">📄 Sullair Argentina</p>', unsafe_allow_html=True)
        st.markdown('<p class="auth-sub">Generador de Certificados</p>', unsafe_allow_html=True)
        tab_login, tab_reg = st.tabs(["🔑 Iniciar sesión", "✏️ Crear cuenta"])
        with tab_login:
            usuario  = st.text_input("Usuario", placeholder="Ej: JICUZA", key="lu").strip().upper()
            password = st.text_input("Contraseña", type="password", key="lp")
            if st.button("Ingresar", key="btn_login"):
                usuarios = cargar_usuarios()
                if usuario in usuarios and usuarios[usuario]["password"] == _hash(password):
                    u = usuarios[usuario]
                    st.session_state.update(logged_in=True, usuario=usuario,
                        nombre=u["nombre"], apellido=u["apellido"],
                        sucursal_usuario=u["sucursal"], mail=u["mail"])
                    st.rerun()
                else:
                    st.error("Usuario o contraseña incorrectos.")
        with tab_reg:
            st.markdown("##### Crear cuenta nueva")
            c1, c2 = st.columns(2)
            with c1: rn = st.text_input("Nombre",   key="rn")
            with c2: ra = st.text_input("Apellido", key="ra")
            rs = st.selectbox("Sucursal", SUCURSALES, key="rs")
            rm = st.text_input("Email", placeholder="nombre@sullair.com", key="rm")
            c3, c4 = st.columns(2)
            with c3: rp1 = st.text_input("Contraseña",         type="password", key="rp1")
            with c4: rp2 = st.text_input("Repetir contraseña", type="password", key="rp2")
            if st.button("Crear cuenta", key="btn_reg"):
                if not all([rn, ra, rm, rp1, rp2]):
                    st.error("Completá todos los campos.")
                elif rp1 != rp2:   st.error("Las contraseñas no coinciden.")
                elif len(rp1) < 6: st.error("Mínimo 6 caracteres.")
                elif "@" not in rm: st.error("Email inválido.")
                else:
                    base = (rn[0] + ra).upper().replace(" ", "")[:10]
                    usuarios = cargar_usuarios()
                    nuevo = base; cnt = 1
                    while nuevo in usuarios:
                        nuevo = f"{base}{cnt}"; cnt += 1
                    usuarios[nuevo] = {"nombre":rn.strip(),"apellido":ra.strip(),
                        "sucursal":rs,"mail":rm.strip(),"password":_hash(rp1)}
                    guardar_usuarios(usuarios)
                    st.success(f"✅ Cuenta creada. Tu usuario es: **{nuevo}**")
                    st.info("Iniciá sesión en la pestaña 'Iniciar sesión'.")
        st.markdown('</div>', unsafe_allow_html=True)

if not st.session_state.get('logged_in'):
    pantalla_auth(); st.stop()

# ─────────────────────────────────────────────────────────────────────────────
# HELPERS GENERALES
# ─────────────────────────────────────────────────────────────────────────────
BASE_DIR  = os.path.dirname(os.path.abspath(__file__))
LOGO_PATH = os.path.join(BASE_DIR, "assets", "logo_sullair.jpg")
VERDE_HEX = "FF71AF47"; BLANCO_HEX = "FFFFFFFF"; NEGRO_HEX = "FF000000"
VERDE_RL  = colors.Color(113/255, 175/255, 71/255)

def med():  return Side(border_style='medium', color=NEGRO_HEX)
def thin(): return Side(border_style='thin',   color=NEGRO_HEX)
def nob():  return Side(border_style=None)

def fmt_fecha(dt):
    if dt is None: return ''
    if isinstance(dt, (datetime, date)): return dt.strftime('%d/%m/%Y')
    return str(dt)

def limpiar(s): return re.sub(r'[\\/*?:"<>|]', '', str(s).strip()).replace('  ', ' ')

def clasificar(un, art, interno):
    un_str = str(un).strip().split()[0] if pd.notna(un) and str(un).strip() else ''
    if un_str in UN_MAP: return UN_MAP[un_str]
    i = str(interno).strip().upper()
    if i.startswith(('E01','A01')): return 'Compresor'
    if i.startswith(('E02','A02')): return 'Generador'
    if i.startswith('E03'):         return 'Plataforma'
    if i.startswith(('E04','A04')): return 'Manipulador'
    if i.startswith('E05'):         return 'Luminaria'
    if i.startswith(('E06','A06')): return 'Movimiento de Suelo'
    if re.match(r'^3\d{3}$', i):   return 'Plataforma'
    if re.match(r'^4\d{3}$', i):   return 'Manipulador'
    FMAP = {'70R':'Generador','120R':'Generador','45KVA':'Generador',
            '375Q':'Compresor','375QH':'Compresor','185Q':'Compresor','3007':'Compresor',
            '860SJ':'Plataforma','1200SJP':'Plataforma',
            'STH1256':'Manipulador','540.170':'Manipulador','4017':'Manipulador',
            'ATABDIS':'Luminaria'}
    for k, v in FMAP.items():
        if k.upper() in str(art).upper(): return v
    return 'Equipo'

def parsear_periodo(raw):
    r = str(raw).strip()
    if re.search(r'\d{4}', r): return r
    return f"{r} {datetime.today().year}"

def leer_vouchers(archivo_bytes):
    df = pd.read_excel(io.BytesIO(archivo_bytes), header=None)
    periodo_raw = str(df.iloc[0, 0]).strip() if pd.notna(df.iloc[0, 0]) else ''
    periodo     = parsear_periodo(periodo_raw)
    headers = [str(v).strip() if pd.notna(v) else f'col_{i}' for i, v in enumerate(df.iloc[1])]
    datos = df.iloc[2:].copy(); datos.columns = headers
    datos = datos.dropna(subset=['Artículo', 'Interno'], how='all')
    return datos, periodo

# ─────────────────────────────────────────────────────────────────────────────
# EXCEL — UNA HOJA POR CERTIFICADO
# ─────────────────────────────────────────────────────────────────────────────
def crear_hoja_xlsx(ws, cert_num, sucursal, direccion, periodo,
                    cliente, articulo, interno, equipo, inicio, fin,
                    cnt_dias, precio, total, moneda, mostrar_precios):

    ws.sheet_view.showGridLines = False
    for col, w in {'A':7.14,'B':4.14,'C':15.43,'D':11.57,'E':8.29,
                   'F':15.0,'G':14.57,'H':12.14,'I':20.57,'J':23.14,
                   'K':6.71,'L':8.43}.items():
        ws.column_dimensions[col].width = w
    for r, h in {3:15.75,4:14.25,5:18.75,6:15.75,7:15.0,
                 8:15.75,9:15.75,10:15.75,11:15.75,12:16.5,
                 13:18.0,14:15.75,17:27.75,18:24.0,19:14.25,
                 20:15.75,21:28.5,22:28.5}.items():
        ws.row_dimensions[r].height = h

    # Borde exterior
    for col in ['B','C','D','E','F','G','H','I','J','K']:
        ws[f'{col}4'].border = Border(top=med(),
            left=med() if col=='B' else nob(), right=med() if col=='K' else nob())
    for row in range(5, 20):
        ws[f'B{row}'].border = Border(left=med())
        ws[f'K{row}'].border = Border(right=med())
    for col in ['B','C','D','E','F','G','H','I','J','K']:
        ws[f'{col}20'].border = Border(bottom=med(),
            left=med() if col=='B' else nob(), right=med() if col=='K' else nob())

    # Logo
    ws.merge_cells('H4:K7')
    if os.path.exists(LOGO_PATH):
        try:
            img = XLImage(LOGO_PATH); img.width=253; img.height=53; img.anchor='I4'
            ws.add_image(img)
        except: pass

    # Título
    ws.merge_cells('C5:E5')
    c=ws['C5']; c.value='CERTIFICACIÓN DE SERVICIOS'
    c.font=Font(name='Calibri',bold=True,size=14); c.alignment=Alignment(horizontal='left')
    ws.merge_cells('C6:D6')
    c=ws['C6']; c.value='Alquiler de equipos'
    c.font=Font(name='Calibri',size=12); c.alignment=Alignment(horizontal='left')

    def lbl(r,col,txt):
        c=ws.cell(r,col,txt); c.font=Font(name='Calibri',bold=True,size=12)
    def val(r,col,txt,h='left',wrap=False):
        c=ws.cell(r,col,txt); c.font=Font(name='Calibri',size=11)
        c.alignment=Alignment(horizontal=h,wrap_text=wrap)

    lbl(8,3,'Certificado N°'); val(8,4,cert_num)
    val(8,9,f'Sucursal {sucursal}',wrap=True)
    lbl(9,3,'Cliente ')
    ws.merge_cells('D9:G9'); val(9,4,cliente,wrap=True)
    val(9,9,direccion,wrap=True)
    lbl(10,3,'Período'); val(10,4,periodo)

    fv=PatternFill('solid',fgColor=VERDE_HEX)
    fhb=Font(name='Calibri',bold=True,size=11,color=BLANCO_HEX)
    ac=Alignment(horizontal='center',vertical='center',wrap_text=True)
    fd=Font(name='Calibri',size=11)
    num_fmt=f'"{moneda}" #,##0.00'

    if mostrar_precios:
        # 8 columnas: C-J
        hdrs=['Equipo','Modelo','Interno','Inicio Periodo','Fin Periodo','Dias','Valor diario (USD)','Total (U$S)']
        cols_n=list(range(3,11))
        hdr_borders=[
            Border(top=med(),bottom=med(),left=med(), right=thin()),
            Border(top=med(),bottom=med(),left=thin(),right=thin()),
            Border(top=med(),bottom=med(),left=thin(),right=thin()),
            Border(top=med(),bottom=med(),left=thin(),right=thin()),
            Border(top=med(),bottom=med(),left=thin(),right=thin()),
            Border(top=med(),bottom=med(),left=thin(),right=thin()),
            Border(top=med(),bottom=med(),left=thin(),right=thin()),
            Border(top=med(),bottom=med(),left=thin(),right=med()),
        ]
        dat_vals=[equipo,articulo.strip(),interno.strip(),
                  fmt_fecha(inicio),fmt_fecha(fin),cnt_dias,precio,total]
        dat_borders=[
            Border(bottom=med(),left=med(), right=thin()),
            Border(bottom=med(),left=thin(),right=thin()),
            Border(bottom=med(),left=thin(),right=thin()),
            Border(bottom=med(),left=thin(),right=thin()),
            Border(bottom=med(),left=thin(),right=thin()),
            Border(bottom=med(),left=thin(),right=thin()),
            Border(bottom=med(),left=thin(),right=thin()),
            Border(bottom=med(),left=thin(),right=med()),
        ]
    else:
        # 6 columnas: C-H (sin precios)
        hdrs=['Equipo','Modelo','Interno','Inicio Periodo','Fin Periodo','Dias']
        cols_n=list(range(3,9))
        hdr_borders=[
            Border(top=med(),bottom=med(),left=med(), right=thin()),
            Border(top=med(),bottom=med(),left=thin(),right=thin()),
            Border(top=med(),bottom=med(),left=thin(),right=thin()),
            Border(top=med(),bottom=med(),left=thin(),right=thin()),
            Border(top=med(),bottom=med(),left=thin(),right=thin()),
            Border(top=med(),bottom=med(),left=thin(),right=med()),
        ]
        dat_vals=[equipo,articulo.strip(),interno.strip(),
                  fmt_fecha(inicio),fmt_fecha(fin),cnt_dias]
        dat_borders=[
            Border(bottom=med(),left=med(), right=thin()),
            Border(bottom=med(),left=thin(),right=thin()),
            Border(bottom=med(),left=thin(),right=thin()),
            Border(bottom=med(),left=thin(),right=thin()),
            Border(bottom=med(),left=thin(),right=thin()),
            Border(bottom=med(),left=thin(),right=med()),
        ]

    # Encabezado tabla
    for col_n, txt, brd in zip(cols_n, hdrs, hdr_borders):
        c=ws.cell(12,col_n,txt); c.fill=fv; c.font=fhb; c.alignment=ac; c.border=brd

    # Datos tabla
    for col_n, v, brd in zip(cols_n, dat_vals, dat_borders):
        c=ws.cell(13,col_n,v); c.font=fd; c.alignment=ac; c.border=brd
        if mostrar_precios and col_n in(9,10): c.number_format=num_fmt

    # Fila 14 — vacía + cant/subtotal
    last_col = cols_n[-1]
    for col_n in cols_n[:-2]:
        c=ws.cell(14,col_n)
        c.border=Border(bottom=med(), left=med() if col_n==cols_n[0] else nob())

    if mostrar_precios:
        c=ws.cell(14,9,'SUBTOTAL'); c.font=Font(name='Calibri',bold=True,size=11)
        c.alignment=Alignment(horizontal='center',vertical='center'); c.border=Border(bottom=med())
        c=ws.cell(14,10,total); c.font=Font(name='Calibri',bold=True,size=11)
        c.number_format=num_fmt; c.alignment=Alignment(horizontal='center',vertical='center')
        c.border=Border(bottom=med(),right=med())
    else:
        # Celda de cant vacía y cierre
        c=ws.cell(14,7); c.border=Border(bottom=med())
        c=ws.cell(14,8,'1 (MES)'); c.font=Font(name='Calibri',bold=True,size=11)
        c.alignment=Alignment(horizontal='center',vertical='center')
        c.border=Border(bottom=med(),right=med())

    # Observaciones
    c=ws.cell(17,3,'OBSERVACIONES:'); c.font=Font(name='Calibri',bold=True,size=12)
    ws.merge_cells('C18:G18'); c=ws['C18']
    c.value='Los valores del presente certificado no incluyen IVA ni otros impuestos'
    c.font=Font(name='Calibri',italic=True,size=11)
    c.alignment=Alignment(horizontal='left',vertical='center',wrap_text=True)
    ws.merge_cells('I18:J19'); c=ws['I18']; c.value='Firma cliente'
    c.font=Font(name='Calibri',size=10)
    c.alignment=Alignment(horizontal='center',vertical='center')

# ─────────────────────────────────────────────────────────────────────────────
# PDF — UN ARCHIVO POR CERTIFICADO (sin CVU, texto ajustado)
# ─────────────────────────────────────────────────────────────────────────────
def crear_pdf(cert_num, sucursal, direccion, periodo, cliente,
              articulo, interno, equipo, inicio, fin,
              cnt_dias, precio, total, moneda, mostrar_precios):

    buf = io.BytesIO()
    PAGE_W, PAGE_H = A4
    c = pdf_canvas.Canvas(buf, pagesize=A4)

    BOX_L=1.75*cm; BOX_R=PAGE_W-1.75*cm
    BOX_T=PAGE_H-2.0*cm; BOX_B=PAGE_H-15.0*cm
    BOX_W=BOX_R-BOX_L

    # Recuadro
    c.setStrokeColor(colors.black); c.setLineWidth(1.2)
    c.rect(BOX_L,BOX_B,BOX_W,BOX_T-BOX_B,stroke=1,fill=0)

    # Logo (arriba derecha, dentro del recuadro)
    LOGO_W=4.6*cm; LOGO_H=1.35*cm
    if os.path.exists(LOGO_PATH):
        c.drawImage(LOGO_PATH,
                    BOX_R-LOGO_W-0.35*cm, BOX_T-LOGO_H-0.55*cm,
                    width=LOGO_W, height=LOGO_H,
                    preserveAspectRatio=True, mask='auto')

    # Título
    TX=BOX_L+0.65*cm
    c.setFont('Helvetica-Bold',11.5); c.setFillColor(colors.black)
    c.drawString(TX, BOX_T-1.15*cm, 'CERTIFICACIÓN DE SERVICIOS')
    c.setFont('Helvetica',9); c.drawString(TX, BOX_T-1.75*cm, 'Alquiler de equipos')

    # Línea separadora
    SEP_Y=BOX_T-2.1*cm; c.setLineWidth(0.4); c.line(BOX_L,SEP_Y,BOX_R,SEP_Y)

    # Encabezado — sin CVU, alineado limpio
    LBL_X=TX; VAL_X=TX+3.0*cm; MID_X=BOX_L+BOX_W*0.52
    RH=0.60*cm; Y0=SEP_Y-0.72*cm

    def draw_lbl(y,txt):
        c.setFont('Helvetica-Bold',9.5); c.setFillColor(colors.black)
        c.drawString(LBL_X,y,txt)

    def draw_val(y,txt,x=None,max_w=None):
        c.setFont('Helvetica',9.5); c.setFillColor(colors.black)
        s=str(txt)
        if max_w:
            # Truncar si es muy largo
            while c.stringWidth(s,'Helvetica',9.5) > max_w and len(s)>1:
                s=s[:-1]
        c.drawString(x or VAL_X, y, s)

    COL_W = MID_X - VAL_X - 0.2*cm   # ancho disponible columna izquierda valor
    MID_W = BOX_R - MID_X - 0.3*cm   # ancho disponible columna derecha

    draw_lbl(Y0,'Certificado N°:');
    draw_val(Y0, str(cert_num))
    draw_val(Y0, f'Sucursal {sucursal}', x=MID_X, max_w=MID_W)

    draw_lbl(Y0-RH,'Cliente:')
    draw_val(Y0-RH, cliente, max_w=COL_W)
    draw_val(Y0-RH, direccion, x=MID_X, max_w=MID_W)

    draw_lbl(Y0-2*RH,'Periodo:')
    draw_val(Y0-2*RH, periodo)

    # Fecha certif (en lugar de CVU)
    c.setFont('Helvetica-Bold',9.5); c.drawString(MID_X, Y0-2*RH, 'Fecha Certif.:')
    c.setFont('Helvetica',9.5)
    c.drawString(MID_X+3.0*cm, Y0-2*RH, datetime.today().strftime('%d/%m/%Y'))

    # Tabla
    TAB_TOP=SEP_Y-3.2*cm; TAB_L=BOX_L+0.35*cm; TAB_W=BOX_W-0.7*cm

    if mostrar_precios:
        CW=[TAB_W*0.22,TAB_W*0.10,TAB_W*0.12,
            TAB_W*0.12,TAB_W*0.12,TAB_W*0.07,
            TAB_W*0.13,TAB_W*0.12]
        hdrs=['Equipo','Modelo','Interno','Desde','Hasta','Dias','Valor diario','Total']
        dat=[equipo, articulo, interno,
             fmt_fecha(inicio), fmt_fecha(fin), cnt_dias,
             f'{moneda} {precio:,.2f}', f'{moneda} {total:,.2f}']
    else:
        CW=[TAB_W*0.28,TAB_W*0.14,TAB_W*0.17,
            TAB_W*0.15,TAB_W*0.15,TAB_W*0.11]
        hdrs=['Equipo','Modelo','Interno','Desde','Hasta','Dias']
        dat=[equipo, articulo, interno,
             fmt_fecha(inicio), fmt_fecha(fin), cnt_dias]

    # Fila vacía de cierre
    data_table = [hdrs, dat, ['']*len(hdrs)]

    t=Table(data_table, colWidths=CW, rowHeights=[18,16,8])
    t.setStyle(TableStyle([
        # Header
        ('BACKGROUND',  (0,0),(-1,0), VERDE_RL),
        ('TEXTCOLOR',   (0,0),(-1,0), colors.white),
        ('FONTNAME',    (0,0),(-1,0), 'Helvetica-Bold'),
        ('FONTSIZE',    (0,0),(-1,0), 8),
        ('ALIGN',       (0,0),(-1,0), 'CENTER'),
        ('VALIGN',      (0,0),(-1,0), 'MIDDLE'),
        ('WORDWRAP',    (0,0),(-1,0), True),
        # Data row
        ('FONTNAME',    (0,1),(-1,1), 'Helvetica'),
        ('FONTSIZE',    (0,1),(-1,1), 8.5),
        ('ALIGN',       (0,1),(-1,1), 'CENTER'),
        ('ALIGN',       (0,1),(0,1),  'LEFT'),
        ('LEFTPADDING', (0,1),(0,1),  4),
        ('VALIGN',      (0,1),(-1,1), 'MIDDLE'),
        ('WORDWRAP',    (0,1),(-1,1), True),
        # Borders
        ('BOX',         (0,0),(-1,-1), 0.8, colors.black),
        ('INNERGRID',   (0,0),(-1,-1), 0.5, colors.black),
    ]))
    t.wrapOn(c, TAB_W, 200)
    t.drawOn(c, TAB_L, TAB_TOP-t._height)

    # Observaciones
    obs_y = TAB_TOP - t._height - 1.2*cm
    c.setFont('Helvetica-Bold',9.5); c.setFillColor(colors.black)
    c.drawString(TX, obs_y, 'OBSERVACIONES:')
    c.setFont('Helvetica-Oblique',8.5)
    c.drawString(TX, obs_y-0.5*cm,
                 'Los valores del presente certificado no incluyen IVA ni otros impuestos')

    # Firma cliente
    c.setFont('Helvetica',8.5)
    c.drawRightString(BOX_R-0.3*cm, obs_y-0.5*cm, 'Firma cliente')

    # Pie fuera del recuadro
    c.setFont('Helvetica',8)
    c.drawRightString(BOX_R, BOX_B-0.55*cm, 'SULLAIR ARGENTINA S.A')

    c.save(); buf.seek(0)
    return buf.read()

# ─────────────────────────────────────────────────────────────────────────────
# PROCESO PRINCIPAL
# ─────────────────────────────────────────────────────────────────────────────
def procesar(archivo_bytes, cert_inicio, periodo_manual, sucursal, direccion,
             mostrar_precios, formato, prog_bar, status_txt):

    datos, periodo_det = leer_vouchers(archivo_bytes)
    periodo = periodo_manual.strip() if periodo_manual.strip() else periodo_det

    cliente_nombre = ''
    if 'Nombre' in datos.columns:
        vals = datos['Nombre'].dropna().unique()
        if len(vals): cliente_nombre = limpiar(str(vals[0]).strip())

    cert_num=cert_inicio; generados=0; log_lines=[]; total_f=len(datos)

    wb = Workbook() if formato=='Excel' else None
    if wb: wb.remove(wb.active)
    zip_buf = io.BytesIO() if formato=='PDF' else None
    zip_file = zipfile.ZipFile(zip_buf,'w',zipfile.ZIP_DEFLATED) if formato=='PDF' else None

    for idx,(_, row) in enumerate(datos.iterrows()):
        art  = str(row.get('Artículo','')).strip()
        int_ = str(row.get('Interno', '')).strip()
        if not art and not int_: continue

        un   = row.get('U.N.','')
        cli  = str(row.get('Nombre','')).strip()
        dias = int(row.get('Cnt Días',0) or 0)
        ini  = row.get('Inicio Período'); fin = row.get('Fin Período')
        prec = float(row.get('Pcio Unitario',0) or 0)
        tot  = float(row.get('Total',0) or 0)
        mon  = str(row.get('Moneda','U$S')).strip().replace('(','').replace(')','').strip()
        eq   = clasificar(un, art, int_)

        if formato == 'Excel':
            nombre_hoja = limpiar(f"{int_[:12].strip()} - {cert_num}")[:31]
            ws = wb.create_sheet(title=nombre_hoja)
            crear_hoja_xlsx(ws,cert_num,sucursal,direccion,periodo,
                            cli,art,int_,eq,ini,fin,dias,prec,tot,mon,mostrar_precios)
        else:
            pdf_bytes = crear_pdf(cert_num,sucursal,direccion,periodo,
                                  cli,art,int_,eq,ini,fin,dias,prec,tot,mon,mostrar_precios)
            # Nombre PDF: Interno-Empresa-Periodo.pdf
            nombre_pdf = limpiar(f"{int_.strip()}-{cliente_nombre}-{periodo}.pdf")
            zip_file.writestr(nombre_pdf, pdf_bytes)

        log_lines.append(f"✅ Cert {cert_num:04d} | {eq:<22} | {int_:<12} | {art}")
        cert_num+=1; generados+=1
        prog_bar.progress(int((idx+1)/total_f*100))
        status_txt.text(f"Generando {idx+1} de {total_f}...")

    nombre_descarga = limpiar(f"Certificación - {cliente_nombre} - {periodo}")

    if formato == 'Excel':
        buf=io.BytesIO(); wb.save(buf); buf.seek(0)
        return {'bytes':buf.read(),'ext':'xlsx','generados':generados,
                'cert_inicio':cert_inicio,'cert_fin':cert_num-1,'periodo':periodo,
                'log':log_lines,'nombre_descarga':f"{nombre_descarga}.xlsx"}
    else:
        zip_file.close(); zip_buf.seek(0)
        return {'bytes':zip_buf.read(),'ext':'zip','generados':generados,
                'cert_inicio':cert_inicio,'cert_fin':cert_num-1,'periodo':periodo,
                'log':log_lines,'nombre_descarga':f"{nombre_descarga}.zip"}

# ─────────────────────────────────────────────────────────────────────────────
# UI PRINCIPAL
# ─────────────────────────────────────────────────────────────────────────────
col_t,col_logout=st.columns([4,1])
with col_t:
    st.markdown('<p class="titulo">📄 Generador de Certificados</p>',unsafe_allow_html=True)
    nc=f'{st.session_state.get("nombre","")} {st.session_state.get("apellido","")}'.strip()
    st.markdown(f'<p class="sub">Sullair Argentina S.A. &nbsp;|&nbsp; 👤 {nc} — {st.session_state.get("sucursal_usuario","")}</p>',unsafe_allow_html=True)
with col_logout:
    st.write("")
    if st.button("🚪 Salir"):
        for k in ['logged_in','usuario','nombre','apellido','sucursal_usuario','mail','resultado']:
            st.session_state.pop(k,None)
        st.rerun()

# Paso 1 — Configuración
st.markdown('<div class="card"><h3>⚙️ Configuración</h3>',unsafe_allow_html=True)
c1,c2=st.columns(2)
with c1: cert_inicio=st.number_input("N° de certificado inicial",min_value=1,value=100,step=1)
with c2: periodo_manual=st.text_input("Período",placeholder="Se detecta automáticamente del Excel")

sucursal_sel=st.selectbox("Sucursal",SUCURSALES,
    index=SUCURSALES.index(st.session_state.get('sucursal_usuario',SUCURSALES[0]))
    if st.session_state.get('sucursal_usuario') in SUCURSALES else 0)
direccion=st.text_input("Dirección de la sucursal",
    placeholder="Ej: Av. San Martín 1234, Neuquén")

c3,c4=st.columns(2)
with c3:
    mostrar_precios=st.toggle("Incluir precios en el certificado",value=True,
        help="Desactivar si el cliente no debe ver los valores económicos")
with c4:
    formato=st.radio("Formato de salida",['Excel','PDF (ZIP)'],horizontal=True)

st.markdown('</div>',unsafe_allow_html=True)

# Paso 2 — Cargar Excel
st.markdown('<div class="card"><h3>📊 Cargar listado de vouchers</h3>',unsafe_allow_html=True)
archivo=st.file_uploader("Arrastrá o seleccioná el Excel de vouchers",
    type=["xlsx","xls"],label_visibility="collapsed")
if archivo:
    st.success(f"✅ **{archivo.name}** cargado")
    datos_prev,periodo_det=leer_vouchers(archivo.read()); archivo.seek(0)
    st.info(f"📅 Período detectado: **{periodo_det}** &nbsp;|&nbsp; 🔢 **{len(datos_prev)}** vouchers")
    with st.expander("👁️ Vista previa — primeros 5 registros"):
        cols_ok=[col for col in ['U.N.','Artículo','Interno','Cnt Días',
                                  'Inicio Período','Fin Período','Pcio Unitario','Total']
                 if col in datos_prev.columns]
        st.dataframe(datos_prev[cols_ok].head(5),use_container_width=True)
st.markdown('</div>',unsafe_allow_html=True)

# Paso 3 — Generar
st.markdown('<div class="card"><h3>⚡ Generar certificados</h3>',unsafe_allow_html=True)
btn_disabled=(archivo is None) or (not direccion.strip())
if archivo is not None and not direccion.strip():
    st.warning("⚠️ Completá la dirección de la sucursal para continuar.")
if st.button("⚡ Generar certificados",disabled=btn_disabled):
    prog=st.progress(0); status=st.empty()
    with st.spinner("Procesando..."):
        resultado=procesar(
            archivo.read(), int(cert_inicio), periodo_manual,
            sucursal_sel, direccion.strip(), mostrar_precios,
            'Excel' if formato=='Excel' else 'PDF',
            prog, status)
    prog.progress(100); status.text("✅ Completado")
    st.session_state["resultado"]=resultado
    with st.expander("📋 Detalle del proceso"):
        for line in resultado["log"]: st.text(line)
st.markdown('</div>',unsafe_allow_html=True)

# Paso 4 — Descarga
if "resultado" in st.session_state:
    r=st.session_state["resultado"]
    st.markdown('<div class="card"><h3>📥 Descargar resultado</h3>',unsafe_allow_html=True)
    c1,c2,c3,c4=st.columns(4)
    with c1: st.markdown(f'<div class="stat"><span class="n">{r["generados"]}</span><span class="l">Certificados</span></div>',unsafe_allow_html=True)
    with c2: st.markdown(f'<div class="stat"><span class="n">{r["cert_inicio"]}</span><span class="l">N° inicial</span></div>',unsafe_allow_html=True)
    with c3: st.markdown(f'<div class="stat"><span class="n">{r["cert_fin"]}</span><span class="l">N° final</span></div>',unsafe_allow_html=True)
    with c4: st.markdown(f'<div class="stat"><span class="n" style="font-size:1rem">{r["periodo"]}</span><span class="l">Período</span></div>',unsafe_allow_html=True)
    st.write("")
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if r["ext"]=="xlsx" else "application/zip"
    icon="📥 Descargar Excel (un certificado por hoja)" if r["ext"]=="xlsx" else "📦 Descargar ZIP (PDFs)"
    st.download_button(label=icon,data=r["bytes"],file_name=r["nombre_descarga"],mime=mime)
    st.markdown('</div>',unsafe_allow_html=True)

st.markdown("---")
st.markdown('<p style="text-align:center;color:#aaa;font-size:.72rem">Sullair Argentina S.A. · Generador Universal de Certificados · v5.1</p>',unsafe_allow_html=True)
